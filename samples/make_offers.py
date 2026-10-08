"""
Make the two sample broker offers for the AI analyst (PDF and Word):

  offers/approve_mukhobola_dairy.*      a well-protected dairy on the edge of the flood plain, priced above the model
  offers/decline_rugunga_rice_mill.*    a rice mill in the deep flood plain with three recent floods, priced far below

Everything in them is fictional: companies, people, phone numbers (+254 700 000 xxx) and emails (example.com).
The people and contact lines are there on purpose, to show the redaction step removing them before the AI reads.
The figures were tuned with the Risk Forge engine so the two verdicts come out as APPROVE and DECLINE.

    python samples/make_offers.py
"""
from pathlib import Path

from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

OUT = Path(__file__).parent / "offers"
PUBLIC = Path(__file__).parent.parent / "web" / "public" / "samples"

BANNER = "SYNTHETIC TEST DOCUMENT - fictional company, people and figures, made for the Risk Forge demo (Kenya Re AI Hackathon)."

APPROVE = {
    "file": "approve_mukhobola_dairy",
    "title": "FACULTATIVE REINSURANCE OFFER - PROPERTY (FLOOD)",
    "sections": [
        ("Offer", [
            "Reference: RF-SAMPLE-A-2027",
            "Cedant: Lakeside General Insurance Ltd (fictional)",
            "Broker: Nzoia Risk Brokers Ltd (fictional)",
            "Reinsurer approached: Kenya Re",
            "Period: 1 January 2027 to 31 December 2027",
        ]),
        ("The insured", [
            "Insured: Mukhobola Fresh Dairy Co-operative Ltd (fictional)",
            "Occupancy: milk collection, pasteurisation and cold storage",
            "Location: Mukhobola village, Budalangi, Busia County",
            "GPS: 0.0823 N, 34.0280 E",
            "The site sits on the eastern edge of the Nzoia flood plain, on ground about 1 m above the surrounding fields.",
        ]),
        ("Buildings", [
            "Building 1 - Processing hall. Construction: reinforced concrete frame with concrete block walls. Floor area: 1,200 m2. Floor level: raised 0.6 m above ground on a concrete plinth. Sum insured: KES 96,000,000. Condition: good.",
            "Building 2 - Cold store. Construction: reinforced concrete frame with insulated panel walls. Floor area: 400 m2. Floor level: raised 0.6 m above ground on a concrete plinth. Sum insured: KES 38,000,000. Condition: good.",
            "Building 3 - Administration block. Construction: brick masonry walls. Floor area: 250 m2. Floor level: plinth raised 0.3 m. Sum insured: KES 14,000,000. Condition: good.",
            "Building 4 - Generator and boiler house. Construction: concrete block walls. Floor area: 80 m2. Floor level: at ground level. Sum insured: KES 5,500,000. Condition: good.",
        ]),
        ("Contents", [
            "Machinery - pasteurisation and packing line, in the processing hall: KES 62,000,000",
            "Stock - finished milk products and packaging, in the cold store: KES 18,000,000",
            "Machinery - standby generators and boiler, in the generator and boiler house: KES 9,000,000",
            "Total sum insured (buildings and contents): KES 242,500,000",
        ]),
        ("Flood history (15 years of records, 2011 to 2025)", [
            "2020 long rains: water reached 0.35 m in the compound; no water entered the raised buildings. Clean-up cost KES 180,000, below the deductible.",
            "No other flood losses in the record.",
        ]),
        ("Risk management", [
            "Flood warning plan linked to the river gauge upstream, with a named duty officer.",
            "Stock kept on pallets; generators and boiler serviced before each rainy season.",
            "Perimeter drainage cleared every March and September.",
        ]),
        ("Terms offered", [
            "Flood deductible: KES 1,000,000 each and every flood event",
            "Flood limit: KES 120,000,000 any one event",
            "Share offered to Kenya Re: 40%",
            "Flood premium (100%): KES 3,400,000",
        ]),
        ("Broker's recommendation", [
            "We recommend acceptance: the main buildings are raised and the site has had no damaging flood in 15 years.",
        ]),
        ("Contacts", [
            "Prepared by: Ms. Achieng Wafula, Senior Broker, Nzoia Risk Brokers Ltd",
            "Tel: +254 700 000 101 | Email: a.wafula@example.com",
            "Risk survey certified by: Eng. Brian Otieno",
            "Insured's contact person: Mr. Peter Barasa, Plant Manager, +254 700 000 102",
        ]),
    ],
}

DECLINE = {
    "file": "decline_rugunga_rice_mill",
    "title": "FACULTATIVE REINSURANCE OFFER - PROPERTY (FLOOD)",
    "sections": [
        ("Offer", [
            "Reference: RF-SAMPLE-B-2027",
            "Cedant: Lakeside General Insurance Ltd (fictional)",
            "Broker: Nzoia Risk Brokers Ltd (fictional)",
            "Reinsurer approached: Kenya Re",
            "Period: 1 January 2027 to 31 December 2027",
        ]),
        ("The insured", [
            "Insured: Bunyala Riverside Rice Millers Ltd (fictional)",
            "Occupancy: paddy storage and rice milling",
            "Location: Rugunga village, Bunyala, Busia County",
            "GPS: 0.0786 N, 34.0010 E",
            "The mill sits on low ground in the Nzoia flood plain, about 600 m from the river bank, below the dyke line.",
        ]),
        ("Buildings", [
            "Building 1 - Rice mill shed. Construction: steel frame with corrugated iron sheet walls and roof. Floor area: 900 m2. Floor level: at ground level. Sum insured: KES 18,000,000. Condition: fair.",
            "Building 2 - Paddy store. Construction: mixed iron sheet and concrete block walls. Floor area: 1,400 m2. Floor level: at ground level. Sum insured: KES 21,000,000. Condition: fair.",
            "Building 3 - Office. Construction: timber frame with mud block walls. Floor area: 120 m2. Floor level: at ground level. Sum insured: KES 2,400,000. Condition: poor (cracked walls after the 2024 flood, not yet repaired).",
            "Building 4 - Drying yard shed. Construction: open corrugated iron sheet shed on steel posts. Floor area: 600 m2. Floor level: at ground level. Sum insured: KES 3,600,000. Condition: fair.",
        ]),
        ("Contents", [
            "Stock - paddy and milled rice, in the paddy store: KES 45,000,000",
            "Machinery - hulling and polishing line, in the rice mill shed: KES 28,000,000",
            "Total sum insured (buildings and contents): KES 118,000,000",
        ]),
        ("Flood history (8 years of records, 2017 to 2024)", [
            "2018: water 1.4 m deep inside the paddy store; settled loss KES 19,400,000.",
            "2020: water 2.1 m deep at the mill after the dyke breach; settled loss KES 27,900,000.",
            "2024: water 1.8 m deep at the mill; settled loss KES 23,600,000.",
        ]),
        ("Risk management", [
            "Sandbags kept on site; no raised storage; no flood warning plan in place.",
        ]),
        ("Terms offered", [
            "Flood deductible: KES 250,000 each and every flood event",
            "Flood limit: KES 80,000,000 any one event",
            "Share offered to Kenya Re: 60%",
            "Flood premium (100%): KES 1,900,000",
        ]),
        ("Broker's recommendation", [
            "The client is a long-standing account and asks for renewal on the expiring terms.",
        ]),
        ("Contacts", [
            "Prepared by: Mr. Kevin Mutua, Account Executive, Nzoia Risk Brokers Ltd",
            "Tel: +254 700 000 201 | Email: k.mutua@example.com",
            "Insured's contact person: Mrs. Esther Nekesa, Managing Director, +254 700 000 202",
        ]),
    ],
}


def pdf(offer: dict, path: Path) -> None:
    styles = getSampleStyleSheet()
    body = styles["BodyText"]
    body.fontSize = 10
    body.leading = 14
    story = [
        Paragraph(f"<font color='#b91c1c' size='8'>{BANNER}</font>", body),
        Spacer(1, 4 * mm),
        Paragraph(offer["title"], styles["Title"]),
    ]
    for head, lines in offer["sections"]:
        story.append(Paragraph(head, styles["Heading3"]))
        for line in lines:
            story.append(Paragraph(line.replace("&", "&amp;").replace("m2", "m²"), body))
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm, bottomMargin=18 * mm,
                            title=offer["title"], author="Risk Forge sample (fictional)")
    doc.build(story)


def docx(offer: dict, path: Path) -> None:
    d = Document()
    d.styles["Normal"].font.size = Pt(10)
    d.add_paragraph(BANNER)
    d.add_heading(offer["title"], level=1)
    for head, lines in offer["sections"]:
        d.add_heading(head, level=2)
        for line in lines:
            d.add_paragraph(line.replace("m2", "m²"))
    d.core_properties.author = "Risk Forge sample (fictional)"
    d.save(str(path))


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    PUBLIC.mkdir(parents=True, exist_ok=True)
    for offer in (APPROVE, DECLINE):
        pdf(offer, OUT / f"{offer['file']}.pdf")
        docx(offer, OUT / f"{offer['file']}.docx")
        # the web app offers the PDFs as "try a sample"
        (PUBLIC / f"{offer['file']}.pdf").write_bytes((OUT / f"{offer['file']}.pdf").read_bytes())
        print("wrote", offer["file"])
